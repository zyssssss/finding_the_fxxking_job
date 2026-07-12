from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


ROOT = Path(r"C:\finding_worker\career-ops")
INPUT = ROOT / "data" / "kdocs-second-pass.json"
LINKS_INPUT = ROOT / "data" / "kdocs-top-shortlist.cleaned.json"
OUTPUT = Path(r"C:\Users\Administrator.DESKTOP-6H09AD4\Downloads\kdocs-second-pass-with-links.xlsx")


with INPUT.open("r", encoding="utf-8") as f:
    payload = json.load(f)

with LINKS_INPUT.open("r", encoding="utf-8") as f:
    links_payload = json.load(f)

links_by_company = {
    item.get("company_name", ""): item
    for item in links_payload.get("shortlist", [])
}

wb = Workbook()
ws_stats = wb.active
ws_stats.title = "stats"
ws_t1 = wb.create_sheet("tier1")
ws_t2 = wb.create_sheet("tier2")
ws_t3 = wb.create_sheet("tier3")
ws_watch = wb.create_sheet("watchlist")
ws_all = wb.create_sheet("all")

header_font = Font(bold=True)
wrap_alignment = Alignment(wrap_text=True, vertical="top")
fills = {
    "tier1": PatternFill(fill_type="solid", fgColor="C6EFCE"),
    "tier2": PatternFill(fill_type="solid", fgColor="FFEB9C"),
    "tier3": PatternFill(fill_type="solid", fgColor="DDEBF7"),
    "watch": PatternFill(fill_type="solid", fgColor="FCE4D6"),
}

columns = [
    "company_name",
    "company_nature",
    "industry",
    "roles",
    "location",
    "second_pass_score",
    "second_pass_tier",
    "recommended",
    "near_match",
    "strong_company_industry",
    "apply_url",
    "official_notice_url",
    "official_careers_url",
    "boss_urls",
    "notes",
    "second_pass_notes",
]

headers = [
    "公司名称",
    "企业性质",
    "行业",
    "岗位/方向",
    "地点",
    "二次评分",
    "二次层级",
    "严格适合",
    "接近适合",
    "公司强行业强",
    "KDocs投递链接",
    "官方公告链接",
    "官网社招链接",
    "Boss链接",
    "初筛备注",
    "二筛备注",
]


def prepare_sheet(ws):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = wrap_alignment
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


for ws in (ws_t1, ws_t2, ws_t3, ws_watch, ws_all):
    prepare_sheet(ws)

ws_stats.append(["metric", "value"])
for cell in ws_stats[1]:
    cell.font = header_font

for key, value in payload.get("stats", {}).items():
    ws_stats.append([key, value])


def item_to_row(item):
    linked_item = links_by_company.get(item.get("company_name", ""), {})
    merged = {
        **item,
        "official_careers_url": linked_item.get("official_careers_url", ""),
        "boss_urls": linked_item.get("boss_urls", []),
    }
    row = []
    for key in columns:
        value = merged.get(key, "")
        if isinstance(value, list):
            value = "\n".join(value)
        row.append(value)
    return row


def append_rows(ws, items, fill):
    for item in items:
        ws.append(item_to_row(item))
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.alignment = wrap_alignment


append_rows(ws_t1, payload.get("tier1", []), fills["tier1"])
append_rows(ws_t2, payload.get("tier2", []), fills["tier2"])
append_rows(ws_t3, payload.get("tier3", []), fills["tier3"])

watch_items = [item for item in payload.get("companies", []) if item.get("second_pass_tier") == "watchlist"]
append_rows(ws_watch, watch_items, fills["watch"])

for item in payload.get("companies", []):
    tier = item.get("second_pass_tier")
    fill = fills["watch"]
    if tier == "tier1":
        fill = fills["tier1"]
    elif tier == "tier2":
        fill = fills["tier2"]
    elif tier == "tier3":
        fill = fills["tier3"]
    ws_all.append(item_to_row(item))
    for cell in ws_all[ws_all.max_row]:
        cell.fill = fill
        cell.alignment = wrap_alignment

widths = {
    "A": 24,
    "B": 16,
    "C": 34,
    "D": 54,
    "E": 24,
    "F": 10,
    "G": 10,
    "H": 10,
    "I": 10,
    "J": 12,
    "K": 42,
    "L": 42,
    "M": 42,
    "N": 52,
    "O": 24,
    "P": 28,
}

for ws in (ws_t1, ws_t2, ws_t3, ws_watch, ws_all):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

ws_stats.column_dimensions["A"].width = 36
ws_stats.column_dimensions["B"].width = 16

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(str(OUTPUT))
