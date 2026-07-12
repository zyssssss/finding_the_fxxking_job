from pathlib import Path
import json, re

ROOT = Path(r"C:\finding_worker\career-ops")
INPUT = ROOT / "data" / "social-jobs-shortlist.json"
OUTPUT_JSON = ROOT / "data" / "social-jobs-shortlist.cleaned.json"
OUTPUT_XLSX = Path(r"C:\Users\Administrator.DESKTOP-6H09AD4\Downloads\social-jobs-shortlist.xlsx")

with INPUT.open("r", encoding="utf-8") as f:
    d = json.load(f)

def extract_keywords(name):
    """从公司名提取关键词片段"""
    name = name.strip()
    parts = []
    # 去掉常见的后缀
    base = re.sub(r'[-—–·\s]*补录$|[-—–·\s]*专场$|[-—–·\s]*管培生.*$|集团|有限公司|股份有限|科技有限|技术有限|进出口有限', '', name)
    # 取前2-4个汉字作为关键词
    if len(base) >= 2:
        parts.append(base[:2])
    if len(base) >= 3:
        parts.append(base[:3])
    if len(base) >= 4:
        parts.append(base[:4])
    # 加上英文名
    eng = re.findall(r'[A-Za-z]{2,}', name)
    parts.extend(eng)
    return [p.lower() for p in parts if len(p) >= 2]

def is_likely_match(url, title, company_name):
    """检查 URL 或标题是否包含公司名关键词"""
    text = (url + ' ' + title).lower()
    keywords = extract_keywords(company_name)
    if not keywords:
        return False
    return any(k in text for k in keywords)

# 额外黑名单：明显不对的域名
blacklist_domains = [
    'linkedin.com', 'indeed.com', 'qcc.com', 'bebee.com', 'nowcoder.com',
    'shushuqiuzhi.com', 'iguopin.com', 'job.mohrss.gov.cn', 'upcv.tech',
    'gaoxiaojob.com', 'goodjobs.cn', 'jrzp.com', 'ncss.cn', 'sohu.com',
    'sina.com.cn', 'lyingjiesheng.com', 'liepin.com', '51job.com',
    'zhaopin.com', 'yingjiesheng.com', '51haojob.com', 'offertoday.com',
    'lingshangdao.com', 'university-hr.com', 'bendibao.com',
    'remote-info.cn', '.edu.cn',
]

cleaned = 0
for item in d["shortlist"]:
    url = item.get("social_careers_url", "")
    title = item.get("social_careers_title", "")
    name = item.get("company_name", "")
    
    if not url:
        continue
    
    # 检查黑名单域名
    try:
        from urllib.parse import urlparse
        host = urlparse(url).hostname or ""
    except:
        host = ""
    
    should_drop = False
    
    for bd in blacklist_domains:
        if bd in host:
            should_drop = True
            break
    
    if not should_drop:
        # 检查是否和公司名匹配
        if not is_likely_match(url, title, name):
            should_drop = True
    
    if should_drop:
        item["social_careers_url"] = ""
        item["social_careers_title"] = ""
        cleaned += 1

d["stats"]["cleaned_official_urls"] = cleaned
d["stats"]["kept_official_urls"] = d["stats"]["with_social"] - cleaned

with OUTPUT_JSON.open("w", encoding="utf-8") as f:
    json.dump(d, f, ensure_ascii=False, indent=2)

# Export XLSX
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

wb = Workbook()
ws_s = wb.active
ws_s.title = "stats"
ws = wb.create_sheet("companies")

hdr = Font(bold=True)
wrp = Alignment(wrap_text=True, vertical="top")
green = PatternFill(fill_type="solid", fgColor="C6EFCE")

cols = ["company_name","company_nature","industry","location","score","kdocs_apply","social_careers_url","social_careers_title","boss_positions"]
heads = ["公司名称","企业性质","行业","地点","评分","KDocs原链接","官网社招链接","社招链接标题","Boss岗位"]

ws_s.append(["metric","value"])
for c in ws_s[1]: c.font = hdr
for k,v in d["stats"].items():
    ws_s.append([k,v])

ws.append(heads)
for c in ws[1]: c.font = hdr

for item in d.get("shortlist", []):
    row = [
        item.get("company_name",""),
        item.get("company_nature",""),
        item.get("industry",""),
        item.get("location",""),
        item.get("score",""),
        item.get("kdocs_apply",""),
        item.get("social_careers_url",""),
        item.get("social_careers_title",""),
        "\n".join([f'{b["title"]}\n{b["url"]}' for b in item.get("boss_positions",[])]),
    ]
    ws.append(row)
    for c in ws[ws.max_row]:
        c.fill = green
        c.alignment = wrp

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
for c,w in {"A":22,"B":14,"C":34,"D":22,"E":8,"F":42,"G":42,"H":32,"I":60}.items():
    ws.column_dimensions[c].width = w
ws_s.column_dimensions["A"].width = 30
ws_s.column_dimensions["B"].width = 14

wb.save(OUTPUT_XLSX)
print(f"Cleaned {cleaned} mis-hits, kept {d['stats']['kept_official_urls']}")
print(f"Saved {OUTPUT_JSON}")
print(f"Saved {OUTPUT_XLSX}")
