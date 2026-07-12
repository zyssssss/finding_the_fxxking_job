from pathlib import Path
import json, re

ROOT = Path(r"C:\finding_worker\career-ops")
INPUT = ROOT / "data" / "social-jobs-shortlist.cleaned.json"
OUTPUT_XLSX = Path(r"C:\Users\Administrator.DESKTOP-6H09AD4\Downloads\social-jobs-final.xlsx")

with INPUT.open("r", encoding="utf-8") as f:
    d = json.load(f)

sz_keywords = ["深圳", "base深圳", "全国出差", "全球出差"]
role_keywords = ["销售", "商务", "渠道", "BD", "bd", "大客户", "客户经理", "解决方案", "售前", "出海", "海外", "国际化", "营销", "管培生", "产品经理"]

def has_shenzhen(loc):
    return any(k in str(loc) for k in sz_keywords)

def role_match_score(roles_text):
    text = str(roles_text).lower()
    return sum(1 for k in role_keywords if k.lower() in text)

for item in d["shortlist"]:
    sz = has_shenzhen(item.get("location", ""))
    has_official = bool(item.get("social_careers_url", ""))
    has_boss = bool(item.get("boss_positions", []))
    role_s = role_match_score(item.get("roles", ""))
    
    priority = 0
    if sz: priority += 5
    if has_official: priority += 3
    if has_boss: priority += 2
    priority += min(role_s, 4)
    
    tier = "P1" if priority >= 10 else "P2" if priority >= 7 else "P3"
    
    item["shenzhen"] = "是" if sz else "否"
    item["link_level"] = ("官网+Boss" if has_official and has_boss else "仅官网" if has_official else "仅Boss" if has_boss else "无")
    item["priority_tier"] = tier
    item["priority_score"] = priority

d["shortlist"].sort(key=lambda x: (-x["priority_score"], -x.get("score", 0)))

# stats
sz_count = sum(1 for x in d["shortlist"] if x["shenzhen"] == "是")
p1 = sum(1 for x in d["shortlist"] if x["priority_tier"] == "P1")
p2 = sum(1 for x in d["shortlist"] if x["priority_tier"] == "P2")
p3 = sum(1 for x in d["shortlist"] if x["priority_tier"] == "P3")

# Export
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

wb = Workbook()
ws_s = wb.active
ws_s.title = "stats"
ws = wb.create_sheet("all")
ws_p1 = wb.create_sheet("P1_优先投")
ws_sz = wb.create_sheet("深圳")

hdr = Font(bold=True)
wrp = Alignment(wrap_text=True, vertical="top")
fills = {
    "P1": PatternFill(fill_type="solid", fgColor="C6EFCE"),
    "P2": PatternFill(fill_type="solid", fgColor="FFEB9C"),
    "P3": PatternFill(fill_type="solid", fgColor="FCE4D6"),
}

cols = ["company_name","company_nature","industry","location","shenzhen","priority_tier","priority_score","score","link_level","social_careers_url","social_careers_title","boss_positions","kdocs_apply"]
heads = ["公司名称","企业性质","行业","地点","深圳","优先级","优先级分","原始分","链接完整度","官网社招链接","社招链接标题","Boss岗位","KDocs原链接"]

ws_s.append(["metric","value"])
for c in ws_s[1]: c.font = hdr
for k,v in {
    "total": len(d["shortlist"]),
    "深圳": sz_count,
    "P1": p1, "P2": p2, "P3": p3,
    "with_social": d["stats"]["with_social"],
    "with_boss": d["stats"]["with_boss"],
    "cleaned_official": d["stats"].get("cleaned_official_urls", 0),
    "kept_official": d["stats"].get("kept_official_urls", 0),
}.items():
    ws_s.append([k,v])

for sheet in (ws, ws_p1, ws_sz):
    sheet.append(heads)
    for c in sheet[1]: c.font = hdr
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

for item in d["shortlist"]:
    row = [
        item.get("company_name",""),
        item.get("company_nature",""),
        item.get("industry",""),
        item.get("location",""),
        item.get("shenzhen",""),
        item.get("priority_tier",""),
        item.get("priority_score",""),
        item.get("score",""),
        item.get("link_level",""),
        item.get("social_careers_url",""),
        item.get("social_careers_title",""),
        "\n".join([f'{b["title"]}\n{b["url"]}' for b in item.get("boss_positions",[])]),
        item.get("kdocs_apply",""),
    ]
    fill = fills.get(item["priority_tier"])
    for sheet in (ws,):
        sheet.append(row)
        for c in sheet[sheet.max_row]:
            c.fill = fill
            c.alignment = wrp
    if item["priority_tier"] == "P1":
        ws_p1.append(row)
        for c in ws_p1[ws_p1.max_row]:
            c.fill = fill
            c.alignment = wrp
    if item["shenzhen"] == "是":
        ws_sz.append(row)
        for c in ws_sz[ws_sz.max_row]:
            c.fill = fill
            c.alignment = wrp

widths = {"A":22,"B":14,"C":34,"D":22,"E":8,"F":10,"G":10,"H":8,"I":14,"J":42,"K":32,"L":60,"M":42}
for sheet in (ws, ws_p1, ws_sz):
    for c,w in widths.items():
        sheet.column_dimensions[c].width = w
ws_s.column_dimensions["A"].width = 30
ws_s.column_dimensions["B"].width = 14

wb.save(OUTPUT_XLSX)
print(f"Total: {len(d['shortlist'])} | 深圳: {sz_count} | P1: {p1} | P2: {p2} | P3: {p3}")
print(str(OUTPUT_XLSX))
