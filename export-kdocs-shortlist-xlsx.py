from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


ROOT = Path(r"C:\finding_worker\career-ops")
INPUT = ROOT / "data" / "kdocs-top-shortlist.cleaned.json"
OUTPUT = Path(r"C:\Users\Administrator.DESKTOP-6H09AD4\Downloads\kdocs-top-shortlist.enhanced.xlsx")


with INPUT.open("r", encoding="utf-8") as f:
    payload = json.load(f)

wb = Workbook()
ws_stats = wb.active
ws_stats.title = "stats"
ws_all = wb.create_sheet("all_shortlist")
ws_sz = wb.create_sheet("shenzhen_priority")
ws_both = wb.create_sheet("official_and_boss")
ws_official_only = wb.create_sheet("official_only")
ws_boss_only = wb.create_sheet("boss_only")
ws_review = wb.create_sheet("needs_review")

header_font = Font(bold=True)
wrap_alignment = Alignment(wrap_text=True, vertical="top")
priority_fills = {
    "P1": PatternFill(fill_type="solid", fgColor="C6EFCE"),
    "P2": PatternFill(fill_type="solid", fgColor="FFEB9C"),
    "P3": PatternFill(fill_type="solid", fgColor="FCE4D6"),
}

stats = payload.get("stats", {})
ws_stats.append(["metric", "value"])
for cell in ws_stats[1]:
    cell.font = header_font

for key, value in stats.items():
    ws_stats.append([key, value])

columns = [
    "company_name",
    "industry",
    "company_nature",
    "roles",
    "location",
    "location_priority",
    "priority_level",
    "score",
    "apply_url",
    "official_notice_url",
    "official_careers_url",
    "official_status",
    "official_query",
    "boss_status",
    "boss_urls",
]

headers = [
    "公司名称",
    "行业",
    "企业性质",
    "岗位/方向",
    "地点",
    "地点优先级",
    "投递优先级",
    "评分",
    "KDocs投递链接",
    "官方公告链接",
    "官网社招链接",
    "官网状态",
    "官网搜索词",
    "Boss状态",
    "Boss链接",
]

def classify_location(location: str) -> str:
    text = str(location or "")
    if "深圳" in text or "base深圳" in text:
        return "high"
    if any(token in text for token in ["广州", "东莞", "珠海", "香港", "全国", "海外", "全球"]):
        return "medium"
    return "low"


def build_row(item: dict) -> list:
    location_priority = classify_location(item.get("location", ""))
    has_official = bool(item.get("official_careers_url"))
    has_boss = bool(item.get("boss_urls"))
    score = int(item.get("score", 0) or 0)

    if location_priority == "high" and has_official and has_boss and score >= 12:
        priority = "P1"
    elif (location_priority in {"high", "medium"}) and (has_official or has_boss) and score >= 10:
        priority = "P2"
    else:
        priority = "P3"

    enriched = {
        **item,
        "location_priority": location_priority,
        "priority_level": priority,
        "official_status": "has_official" if has_official else "missing_official",
        "boss_status": "has_boss" if has_boss else "missing_boss",
    }

    row = []
    for key in columns:
        value = enriched.get(key, "")
        if isinstance(value, list):
            value = "\n".join(value)
        row.append(value)
    return row


def prepare_sheet(ws):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = wrap_alignment


for sheet in (ws_all, ws_sz, ws_both, ws_official_only, ws_boss_only, ws_review):
    prepare_sheet(sheet)

for item in payload.get("shortlist", []):
    row = build_row(item)
    ws_all.append(row)

    location_priority = row[5]
    priority_level = row[6]
    has_official = bool(row[10])
    has_boss = bool(row[13])

    if location_priority == "high":
        ws_sz.append(row)
    if has_official and has_boss:
        ws_both.append(row)
    elif has_official and not has_boss:
        ws_official_only.append(row)
    elif has_boss and not has_official:
        ws_boss_only.append(row)
    else:
        ws_review.append(row)

widths = {
    "A": 24,
    "B": 32,
    "C": 14,
    "D": 50,
    "E": 24,
    "F": 12,
    "G": 10,
    "H": 10,
    "I": 40,
    "J": 40,
    "K": 40,
    "L": 18,
    "M": 28,
    "N": 18,
    "O": 50,
}

data_sheets = (ws_all, ws_sz, ws_both, ws_official_only, ws_boss_only, ws_review)

for ws in (ws_stats, *data_sheets):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = wrap_alignment

for col, width in widths.items():
    for ws in data_sheets:
        ws.column_dimensions[col].width = width

ws_stats.column_dimensions["A"].width = 28
ws_stats.column_dimensions["B"].width = 16

for ws in data_sheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

for ws in data_sheets:
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        priority_cell = row[6]
        fill = priority_fills.get(priority_cell.value)
        if fill:
            for cell in row:
                cell.fill = fill

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(str(OUTPUT))
