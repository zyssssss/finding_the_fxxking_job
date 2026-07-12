from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


ROOT = Path(r"C:\finding_worker\career-ops")
INPUT = ROOT / "data" / "kdocs-full.filtered.json"
OUTPUT = Path(r"C:\Users\Administrator.DESKTOP-6H09AD4\Downloads\kdocs-full.filtered.enhanced.xlsx")


with INPUT.open("r", encoding="utf-8") as f:
    payload = json.load(f)

wb = Workbook()
ws_stats = wb.active
ws_stats.title = "stats"
ws_recommended = wb.create_sheet("recommended")
ws_near = wb.create_sheet("near_match")
ws_strong = wb.create_sheet("strong_company_industry")
ws_all = wb.create_sheet("all_companies")

header_font = Font(bold=True)
wrap_alignment = Alignment(wrap_text=True, vertical="top")

fill_map = {
    "recommended": PatternFill(fill_type="solid", fgColor="C6EFCE"),
    "near_match": PatternFill(fill_type="solid", fgColor="FFEB9C"),
    "strong": PatternFill(fill_type="solid", fgColor="DDEBF7"),
}

columns = [
    "company_name",
    "industry",
    "company_nature",
    "roles",
    "location",
    "score",
    "recommended",
    "near_match",
    "strong_company_industry",
    "industry_class",
    "location_class",
    "apply_url",
    "official_notice_url",
    "notes",
]

headers = [
    "公司名称",
    "行业",
    "企业性质",
    "岗位/方向",
    "地点",
    "评分",
    "严格适合",
    "接近适合",
    "公司强行业强",
    "行业分类",
    "地点分类",
    "KDocs投递链接",
    "官方公告链接",
    "备注",
]


def prepare_sheet(ws):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = wrap_alignment
    ws.freeze_panes = "A2"


def item_to_row(item):
    row = []
    for key in columns:
      value = item.get(key, "")
      if isinstance(value, list):
          value = "\n".join(value)
      row.append(value)
    return row


prepare_sheet(ws_recommended)
prepare_sheet(ws_near)
prepare_sheet(ws_strong)
prepare_sheet(ws_all)

ws_stats.append(["metric", "value"])
for cell in ws_stats[1]:
    cell.font = header_font

for key, value in payload.get("stats", {}).items():
    ws_stats.append([key, value])


def append_rows(ws, items, fill=None):
    for item in items:
        ws.append(item_to_row(item))
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
                cell.alignment = wrap_alignment


append_rows(ws_recommended, payload.get("recommended", []), fill_map["recommended"])
append_rows(ws_near, payload.get("near_match", []), fill_map["near_match"])
append_rows(ws_strong, payload.get("strong_company_industry", []), fill_map["strong"])

for item in payload.get("companies", []):
    if item.get("recommended"):
        fill = fill_map["recommended"]
    elif item.get("near_match"):
        fill = fill_map["near_match"]
    elif item.get("strong_company_industry"):
        fill = fill_map["strong"]
    else:
        fill = None
    ws_all.append(item_to_row(item))
    if fill:
        for cell in ws_all[ws_all.max_row]:
            cell.fill = fill

for ws in (ws_recommended, ws_near, ws_strong, ws_all):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = wrap_alignment
    ws.auto_filter.ref = ws.dimensions

widths = {
    "A": 24,
    "B": 34,
    "C": 14,
    "D": 52,
    "E": 26,
    "F": 8,
    "G": 10,
    "H": 10,
    "I": 12,
    "J": 14,
    "K": 12,
    "L": 42,
    "M": 42,
    "N": 36,
}

for ws in (ws_recommended, ws_near, ws_strong, ws_all):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

ws_stats.column_dimensions["A"].width = 32
ws_stats.column_dimensions["B"].width = 16

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(str(OUTPUT))
