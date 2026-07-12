from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = Path(r"C:\finding_worker\career-ops")
INPUT = ROOT / "data" / "social-jobs-shortlist.json"
OUTPUT = Path(r"C:\Users\Administrator.DESKTOP-6H09AD4\Downloads\social-jobs-shortlist.xlsx")

with INPUT.open("r", encoding="utf-8") as f:
    d = json.load(f)

wb = Workbook()
ws_s = wb.active
ws_s.title = "stats"
ws = wb.create_sheet("companies")

hdr = Font(bold=True)
wrp = Alignment(wrap_text=True, vertical="top")
green = PatternFill(fill_type="solid", fgColor="C6EFCE")

cols = ["company_name","company_nature","industry","location","score","kdocs_apply","social_careers_url","social_careers_title","boss_positions"]
heads = ["公司名称","企业性质","行业","地点","评分","KDocs原链接","官网社招链接","社招链接标题","Boss岗位"]

ws_s.append(["metric","value"])
for c in ws_s[1]: c.font = hdr
for k,v in d["stats"].items():
    ws_s.append([k,v])

ws.append(heads)
for c in ws[1]: c.font = hdr

for item in d.get("shortlist", []):
    row = [
        item.get("company_name",""),
        item.get("company_nature",""),
        item.get("industry",""),
        item.get("location",""),
        item.get("score",""),
        item.get("kdocs_apply",""),
        item.get("social_careers_url",""),
        item.get("social_careers_title",""),
        "\n".join([f'{b["title"]}\n{b["url"]}' for b in item.get("boss_positions",[])]),
    ]
    ws.append(row)
    for c in ws[ws.max_row]:
        c.fill = green
        c.alignment = wrp

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
for c,w in {"A":22,"B":14,"C":34,"D":22,"E":8,"F":42,"G":42,"H":32,"I":60}.items():
    ws.column_dimensions[c].width = w
ws_s.column_dimensions["A"].width = 30
ws_s.column_dimensions["B"].width = 14

wb.save(OUTPUT)
print(str(OUTPUT))
